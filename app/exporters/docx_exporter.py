"""
Módulo: docx_exporter.py.
Responsabilidade: exportar estrutura e formatação Markdown para DOCX.
Autor: Open Doc Contributors
Criado em: 2026
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt
from markdown_it.token import Token

from app.exporters.markdown_document import tokens

BODY_SIZE = 11
CODE_SIZE = 10


def _inline(paragraph: object, children: list[Token]) -> None:
    bold = False
    italic = False
    link = ""
    for token in children:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type == "link_open":
            link = token.attrGet("href") or ""
        elif token.type == "link_close":
            if link:
                paragraph.add_run(" (" + link + ")")
            link = ""
        elif token.type in {"text", "code_inline", "image"}:
            content = token.content
            if token.type == "image":
                content = "[Imagem: " + content + "]"
            run = paragraph.add_run(content)
            run.bold = bold
            run.italic = italic
            if token.type == "code_inline":
                run.font.name = "Courier New"
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run("\n")


def export_docx(markdown: str, target: Path, title: str) -> None:
    """Gera DOCX com títulos, listas, tabelas, citações e trechos de código."""
    document = Document()
    document.core_properties.title = title
    document.styles["Normal"].font.size = Pt(BODY_SIZE)
    stream = tokens(markdown)
    heading = None
    lists: list[str] = []
    quoted = False
    paragraph = None
    table = None
    row = None
    column = 0
    for token in stream:
        kind = token.type
        if kind == "heading_open":
            heading = int(token.tag[1:])
        elif kind == "heading_close":
            heading = None
        elif kind == "bullet_list_open":
            lists.append("List Bullet")
        elif kind == "ordered_list_open":
            lists.append("List Number")
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "blockquote_open":
            quoted = True
        elif kind == "blockquote_close":
            quoted = False
        elif kind == "table_open":
            table = document.add_table(rows=0, cols=0)
            table.style = "Table Grid"
        elif kind == "table_close":
            table = None
        elif kind == "tr_open":
            row = None
            column = 0
        elif kind in {"th_open", "td_open"}:
            if len(table.columns) <= column:
                from docx.shared import Inches

                table.add_column(Inches(1))
            if row is None:
                row = table.add_row()
            paragraph = row.cells[column].paragraphs[0]
            column += 1
        elif kind == "paragraph_open":
            style = lists[-1] if lists else "Quote" if quoted else None
            paragraph = document.add_paragraph(style=style)
        elif kind == "inline":
            if heading:
                paragraph = document.add_heading(level=heading)
            if paragraph is None:
                paragraph = document.add_paragraph()
            _inline(paragraph, token.children or [])
        elif kind in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(token.content.rstrip())
            run.font.name = "Courier New"
            run.font.size = Pt(CODE_SIZE)
        elif kind == "hr":
            document.add_paragraph("—")
    document.save(str(target))
