from zipfile import ZipFile
from docx import Document
from apps.open_doc.exporters.docx_exporter import export_docx
from apps.open_doc.parsers.docx.package import DocxPackage
from apps.open_doc.parsers.docx.namespaces import NS
from apps.open_doc.parsers.docx.relationships import relationships
from test_docx_import import imported


def test_unknown_parts_and_relationships(imported, tmp_path):
    e, note, source = imported
    target = tmp_path / "export.docx"
    note["content"]["blocks"][0]["children"][0]["text"] = "Título editado"
    export_docx(
        note["content"],
        target,
        note["title"],
        package_directory=e.package_directory(note["id"]),
    )
    assert Document(target).paragraphs[0].text == "Título editado"
    with ZipFile(source) as a, ZipFile(target) as b:
        assert b.testzip() is None and len(b.namelist()) == len(set(b.namelist()))
        assert a.read("customXml/item1.xml") == b.read("customXml/item1.xml")
        for name in a.namelist():
            if name not in {
                "word/document.xml",
                "word/numbering.xml",
                "word/_rels/document.xml.rels",
            }:
                assert a.read(name) == b.read(name), name
    package = DocxPackage.open(target)
    ids = {r.get("Id") for r in package.xml("word/_rels/document.xml.rels")}
    for n in package.xml("word/document.xml").iter():
        for key, val in n.attrib.items():
            if key.startswith("{" + NS["r"] + "}"):
                assert val in ids
    for part in package.parts:
        if part.endswith(".xml"):
            for rel in relationships(package, part).values():
                if not rel["external"]:
                    assert rel["target"] in package.parts


def test_reread_structure(imported, tmp_path):
    e, note, _ = imported
    target = tmp_path / "roundtrip.docx"
    export_docx(
        note["content"],
        target,
        note["title"],
        package_directory=e.package_directory(note["id"]),
    )
    m = e.import_docx(target)["content"]
    assert m["sections"][1]["page_width_pt"] == 792
    t = next(n for n in m["blocks"] if n["type"] == "table")
    assert (
        t["rows"][0]["cells"][0]["col_span"] == 2
        and t["rows"][1]["cells"][2]["row_span"] == 2
    )
    assert any(
        r.get("style", {}).get("color") == "FF0000"
        for p in m["blocks"]
        for r in p.get("children", [])
    )


def test_safe_html(imported, tmp_path):
    from apps.open_doc.exporters.html_exporter import export_html

    e, note, _ = imported
    target = tmp_path / "rich.html"
    note["content"]["blocks"][0]["children"][0]["text"] = "<script>bad()</script>"
    export_html(
        note["content"],
        target,
        note["title"],
        package_directory=e.package_directory(note["id"]),
    )
    text = target.read_text()
    assert (
        "<script>" not in text
        and "&lt;script&gt;" in text
        and "font-family" in text
        and "<table>" in text
        and "data:image/png;base64," in text
    )


def test_native_package(tmp_path):
    from test_document_model import model

    m = model()
    m["blocks"][0]["list_info"] = {"level": 0, "format": "decimal"}
    target = tmp_path / "native.docx"
    export_docx(m, target, "Novo")
    assert Document(target).paragraphs[0].text == "Preservar"
    assert (
        DocxPackage.open(target).xml("word/document.xml").find(".//w:numPr", NS)
        is not None
    )
