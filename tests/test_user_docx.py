"""Optional private regression file; never redistribute the user's academic DOCX."""

import os
from pathlib import Path
from zipfile import ZipFile
from docx import Document
import pytest
from app.config import Config
from app.editor import Editor
from app.exporters.docx_exporter import export_docx


def test_user_docx_text_fonts_footer_and_geometry(tmp_path):
    source = Path(
        os.environ.get(
            "OPEN_DOC_USER_FIXTURE",
            str(
                Path(__file__).resolve().parents[2]
                / "docx-input"
                / "Lista_Exercicios_Fixacao_Aula_02_ML.docx"
            ),
        )
    )
    if not source.exists():
        pytest.skip("Private user DOCX not supplied; set OPEN_DOC_USER_FIXTURE.")
    original = source.read_bytes()
    word = Document(source)
    editor = Editor(Config(tmp_path / "workspace"))
    note = editor.import_docx(source)
    model = note["content"]
    declared = {r.font.name for p in word.paragraphs for r in p.runs if r.font.name}
    actual = {
        r.get("style", {}).get("font_family")
        for p in model["blocks"]
        for r in p.get("children", [])
    }
    assert declared <= actual and model["footers"]
    first = model["sections"][0]
    section = word.sections[0]
    assert (
        first["page_width_pt"] == section.page_width.pt
        and first["page_height_pt"] == section.page_height.pt
    )
    for side in ("left", "right", "top", "bottom"):
        assert first["margin_" + side + "_pt"] == getattr(section, side + "_margin").pt
    editor.save(
        note["id"], note["title"], "", note["revision"], model, "opendoc-rich-v1"
    )
    saved = editor.get(note["id"])
    assert saved["markdown"]
    target = tmp_path / "roundtrip-user.docx"
    export_docx(
        saved["content"],
        target,
        note["title"],
        package_directory=editor.package_directory(note["id"]),
    )
    assert [p.text for p in Document(target).paragraphs] == [
        p.text for p in word.paragraphs
    ]
    with ZipFile(source) as a, ZipFile(target) as b:
        assert b.testzip() is None
        for name in a.namelist():
            if name.startswith("word/footer"):
                assert a.read(name) == b.read(name)
    assert source.read_bytes() == original
    assert Editor(editor.config).current()["content"] == model
