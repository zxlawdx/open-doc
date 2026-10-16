"""Synthetic DOCX with typography, notes, lists, merged tables, images and sections."""

from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from apps.open_doc.parsers.docx.namespaces import W, R, REL, w


def fixture_docx(target):
    doc = Document()
    doc.add_heading("Documento de validação", 1)
    p = doc.add_paragraph()
    for name, size in [
        ("Arial", 10),
        ("Times New Roman", 12),
        ("Courier New", 9),
        ("Aptos", 11),
    ]:
        r = p.add_run(name + " ")
        r.font.name = name
        r.font.size = Pt(size)
    r = p.add_run("Formatado")
    r.bold = True
    r.italic = True
    r.underline = True
    r.font.strike = True
    r.font.color.rgb = RGBColor(255, 0, 0)
    e = OxmlElement("w:highlight")
    e.set(qn("w:val"), "yellow")
    r._r.get_or_add_rPr().append(e)
    for key in ("superscript", "subscript"):
        r = p.add_run(key)
        setattr(r.font, key, True)
    p.alignment = 1
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    for level, fmt in [(0, "decimal"), (1, "lowerLetter"), (2, "lowerRoman")]:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(40 + level))
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        for tag, value in [
            ("start", "1"),
            ("numFmt", fmt),
            ("lvlText", "%" + str(level + 1) + ")"),
        ]:
            e = OxmlElement("w:" + tag)
            e.set(qn("w:val"), value)
            lvl.append(e)
        abstract.append(lvl)
        doc.part.numbering_part.element.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(40 + level))
        e = OxmlElement("w:abstractNumId")
        e.set(qn("w:val"), str(40 + level))
        num.append(e)
        doc.part.numbering_part.element.append(num)
        p = doc.add_paragraph("Nível " + str(level))
        np = OxmlElement("w:numPr")
        for tag, value in [("ilvl", str(level)), ("numId", str(40 + level))]:
            e = OxmlElement("w:" + tag)
            e.set(qn("w:val"), value)
            np.append(e)
        p._p.get_or_add_pPr().append(np)
    doc.add_paragraph("Marcador", style="List Bullet")
    table = doc.add_table(rows=3, cols=4)
    for i, row in enumerate(table.rows):
        for j, c in enumerate(row.cells):
            c.text = f"R{i}C{j}"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(0, 0).paragraphs[0].runs[0].bold = True
    table.cell(1, 0).add_paragraph("Lista na célula", style="List Bullet")
    e = OxmlElement("w:shd")
    e.set(qn("w:fill"), "FFFF00")
    table.cell(0, 0)._tc.get_or_add_tcPr().append(e)
    image = target.with_suffix(".png")
    Image.new("RGB", (30, 20), "red").save(image)
    doc.add_picture(str(image))
    doc.add_page_break()
    p = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    rid = doc.part.relate_to("https://example.com", R + "/hyperlink", is_external=True)
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Link seguro"
    run.append(text)
    link.append(run)
    p._p.append(link)
    opaque = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    op = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Opaque preservado"
    r.append(t)
    op.append(r)
    content.append(op)
    opaque.append(content)
    doc._element.body.insert(1, opaque)
    doc.sections[0].header.paragraphs[0].text = "Cabeçalho original"
    doc.sections[0].footer.paragraphs[0].text = "Rodapé original"
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width = Pt(792)
    section.page_height = Pt(612)
    section.left_margin = Pt(54)
    doc.add_paragraph("Segunda seção")
    doc.save(target)
    with ZipFile(target) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    p = root.find(w("body")).findall(w("p"))[-1]
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    types = etree.fromstring(parts["[Content_Types].xml"])
    for kind, tag in [("footnotes", "footnote"), ("endnotes", "endnote")]:
        r = etree.SubElement(p, w("r"))
        e = etree.SubElement(r, w(tag + "Reference"))
        e.set(w("id"), "1")
        notes = etree.Element(w(kind), nsmap={"w": W})
        entry = etree.SubElement(notes, w(tag))
        entry.set(w("id"), "1")
        text = etree.SubElement(
            etree.SubElement(etree.SubElement(entry, w("p")), w("r")), w("t")
        )
        text.text = "Nota " + kind
        parts["word/" + kind + ".xml"] = etree.tostring(notes)
        etree.SubElement(
            rels,
            "{" + REL + "}Relationship",
            Id="rId" + kind,
            Type=R + "/" + kind,
            Target=kind + ".xml",
        )
        etree.SubElement(
            types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            PartName="/word/" + kind + ".xml",
            ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml."
            + kind
            + "+xml",
        )
    for tag, value in [
        ("fldChar", "begin"),
        ("instrText", " PAGE "),
        ("fldChar", "separate"),
        ("t", "7"),
        ("fldChar", "end"),
    ]:
        r = etree.SubElement(p, w("r"))
        e = etree.SubElement(r, w(tag))
        if tag == "fldChar":
            e.set(w("fldCharType"), value)
        else:
            e.text = value
    parts["word/document.xml"] = etree.tostring(root)
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
    parts["[Content_Types].xml"] = etree.tostring(types)
    parts["customXml/item1.xml"] = b"<custom>Keep original bytes</custom>"
    with ZipFile(target, "w", ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return target
