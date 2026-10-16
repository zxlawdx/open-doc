"""Destinos nativos, gravação atômica, catálogo local e operações de documento."""

from pathlib import Path
from unittest.mock import Mock
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfReader

from apps.open_doc.bridge import OpenDocBridge
from apps.open_doc.services.container import configure
from apps.open_doc.services.preferences_service import Config
from apps.open_doc.services.export_service import ExportError
from apps.open_doc.services.font_service import FontService


@pytest.fixture
def runtime(tmp_path):
    return configure(Config(tmp_path / "workspace"))


@pytest.mark.parametrize("extension", ["docx", "pdf", "html", "md"])
def test_export_save_dialog(runtime, tmp_path, extension):
    import webview

    note = runtime.editor.create("Acentuação e espaços", "# Olá\n\nTexto **forte**.")
    destination = tmp_path / "Área de Trabalho" / "Nome escolhido.outro"
    destination.parent.mkdir()
    bridge = OpenDocBridge()
    bridge._window = Mock()
    bridge._window.create_file_dialog.return_value = (str(destination),)
    result = bridge.export_document(note["id"], note["revision"], extension)
    assert "error" not in result, result
    target = destination.with_suffix("." + extension)
    assert Path(result["path"]) == target and target.stat().st_size > 0
    assert (
        bridge._window.create_file_dialog.call_args.args[0] == webview.FileDialog.SAVE
    )
    if extension == "docx":
        assert Document(target).paragraphs
        with ZipFile(target) as archive:
            assert archive.testzip() is None
    if extension == "pdf":
        assert PdfReader(target).pages
    if extension == "html":
        assert "<!doctype html>" in target.read_text()
    if extension == "md":
        assert "**forte**" in target.read_text()
    assert not list(target.parent.glob(".opendoc-*"))


def test_export_cancel(runtime, tmp_path):
    bridge = OpenDocBridge()
    bridge._window = Mock()
    bridge._window.create_file_dialog.return_value = None
    note = runtime.editor.current()
    assert bridge.export_document(note["id"], note["revision"], "docx") == {
        "cancelled": True
    }
    assert not list(runtime.exports.default_export_directory().iterdir())


def test_export_exception_returns_safe_error(runtime, tmp_path, monkeypatch):
    from apps.open_doc.exporters import pdf_exporter

    target = tmp_path / "antigo.pdf"
    target.write_bytes(b"previous")

    def fail(*args, **kwargs):
        args[1].write_bytes(b"incomplete")
        raise RuntimeError("private internal failure")

    monkeypatch.setattr(pdf_exporter, "export_pdf", fail)
    note = runtime.editor.current()
    with pytest.raises(ExportError):
        runtime.exports.export(note["id"], note["revision"], "pdf", target)
    assert target.read_bytes() == b"previous"
    bridge = OpenDocBridge()
    bridge._window = Mock()
    bridge._window.create_file_dialog.return_value = (str(target),)
    assert (
        "private"
        not in bridge.export_document(note["id"], note["revision"], "pdf")["error"]
    )
    assert not list(tmp_path.glob(".opendoc-*"))


def make_font(path, ttf):
    """Minimal original font fixture; no platform fonts or binary redistribution."""
    from fontTools.fontBuilder import FontBuilder

    builder = FontBuilder(1000, isTTF=ttf)
    builder.setupGlyphOrder([".notdef", "space"])
    builder.setupCharacterMap({32: "space"})
    builder.setupHorizontalMetrics({".notdef": (600, 0), "space": (300, 0)})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable(
        {
            "familyName": "OpenDoc Test " + ("TTF" if ttf else "OTF"),
            "styleName": "Regular",
            "uniqueFontIdentifier": "OpenDocTest" + str(ttf),
            "fullName": "OpenDoc Test",
            "psName": "OpenDocTest" + str(ttf),
        }
    )
    builder.setupOS2(
        sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200
    )
    builder.setupPost()
    if ttf:
        from fontTools.pens.ttGlyphPen import TTGlyphPen

        builder.setupGlyf(
            {name: TTGlyphPen(None).glyph() for name in [".notdef", "space"]}
        )
        builder.setupMaxp()
    else:
        from fontTools.pens.t2CharStringPen import T2CharStringPen

        builder.setupCFF(
            "OpenDocTestOTF",
            {},
            {
                name: T2CharStringPen(600, None).getCharString()
                for name in [".notdef", "space"]
            },
            {},
        )
    builder.save(path)


@pytest.mark.parametrize("extension", ["ttf", "otf"])
def test_import_font_and_persistence(runtime, tmp_path, extension):
    path = tmp_path / ("Fonte local." + extension)
    make_font(path, extension == "ttf")
    catalog = runtime.fonts.import_files([path])
    entry = catalog["imported"][0]
    assert entry["family"] in catalog["fonts"]
    restored = FontService(runtime.config)
    assert restored.resolve(entry["family"]).is_file()
    assert len(restored.import_files([path])["imported"]) == 1
    assert restored.remove(entry["id"])["imported"] == []
    assert path.exists()
    with pytest.raises(ValueError):
        restored.remove("system-font")


def test_invalid_font(runtime, tmp_path):
    path = tmp_path / "bad.ttf"
    path.write_bytes(b"not a font")
    with pytest.raises(Exception):
        runtime.fonts.import_files([path])
    assert runtime.fonts.catalog()["imported"] == []


def test_document_rename_duplicate_delete(runtime):
    note = runtime.editor.create("Original", "Conteúdo")
    renamed = runtime.editor.rename(note["id"], "Renomeado", note["revision"])
    assert renamed["title"] == "Renomeado"
    copy = runtime.editor.duplicate(note["id"])
    assert copy["id"] != note["id"] and copy["markdown"] == "Conteúdo"
    runtime.editor.delete(note["id"], renamed["revision"])
    with pytest.raises(ValueError):
        runtime.editor.get(note["id"])
    assert runtime.editor.get(copy["id"])["markdown"] == "Conteúdo"


def test_page_break_setup_docx_roundtrip(runtime, tmp_path):
    from apps.open_doc.document_model.schema import RichDocument

    model = RichDocument().to_dict()
    model["blocks"] = [
        {"id": "break-1", "type": "page_break", "manual": True},
        {
            "id": "p-1",
            "type": "paragraph",
            "style": {},
            "children": [
                {
                    "id": "run-1",
                    "type": "text",
                    "text": "Após quebra",
                    "style": {"font_family": "Fonte preservada"},
                }
            ],
        },
    ]
    model["sections"] = [
        {
            "id": "section-1",
            "page_width_pt": 792,
            "page_height_pt": 612,
            "orientation": "landscape",
            "margin_top_pt": 30,
            "margin_bottom_pt": 40,
            "margin_left_pt": 50,
            "margin_right_pt": 60,
        }
    ]
    note = runtime.editor.create_rich_copy("Páginas", model)
    target = tmp_path / "Páginas.docx"
    runtime.exports.export(note["id"], note["revision"], "docx", target)
    doc = Document(target)
    assert 'w:type="page"' in doc.element.xml
    assert doc.sections[0].page_width.pt == 792
    assert doc.sections[0].left_margin.pt == 50
    reopened = runtime.editor.import_docx(target)
    assert reopened["content"]["sections"][0]["orientation"] == "landscape"
    assert "Fonte preservada" in str(reopened["content"])


@pytest.mark.parametrize("suffix", [".md", ".txt"])
def test_bridge_import_text_opens_editable_document(runtime, tmp_path, suffix):
    content = "Olá, ação e acentuação 🐧\n" * 6000
    source = tmp_path / ("Minha nota" + suffix)
    source.write_text(content, encoding="utf-8")
    bridge = OpenDocBridge()
    bridge._window = Mock()
    bridge._window.create_file_dialog.return_value = (str(source),)
    result = bridge.import_document()
    assert result.get("opened") is True and result.get("document"), result
    assert (
        result["document"]["title"] == "Minha nota"
        and result["document"]["markdown"] == content
    )
    assert (
        runtime.fixed_documents.active_kind() == "rich"
        and source.read_text(encoding="utf-8") == content
    )
    imported = runtime.organizer.list_files()
    assert imported and imported[0]["capabilities"]["open"]
