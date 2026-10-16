"""
Módulo: test_editor.py.
Responsabilidade: verificar persistência, revisões, recuperação e exportações.
Autor: Open Doc Contributors
Criado em: 2026
"""

from __future__ import annotations

from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

import pytest
from docx import Document
from pypdf import PdfReader

from apps.open_doc.services.preferences_service import Config
from apps.open_doc.services.document_service import Editor, RevisionConflict
from apps.open_doc.exporters.docx_exporter import export_docx
from apps.open_doc.exporters.pdf_exporter import export_pdf

MARKDOWN = (
    "# Ciência e ação\n\n"
    "Português **forte**, *itálico*, [link](https://example.com).\n\n"
    "- Primeiro\n- Segundo\n\n"
    "> Uma citação importante.\n\n"
    "\x60\x60\x60python\nprint('Olá')\n\x60\x60\x60\n\n"
    "| Nome | Valor |\n| --- | --- |\n| Estudo | 42 |\n"
)


def test_session_and_unicode(tmp_path: Path) -> None:
    """Restaura conteúdo Unicode e documento ativo após reinicialização."""
    config = Config(tmp_path)
    editor = Editor(config)
    note = editor.create("Ciência", MARKDOWN)
    saved = editor.save(note["id"], "Novo título", MARKDOWN + "Fim", note["revision"])
    restored = Editor(config).current()
    assert restored["id"] == note["id"]
    assert restored["markdown"] == MARKDOWN + "Fim"
    assert restored["revision"] == saved["revision"]


def test_conflict_preserves_newer_content(tmp_path: Path) -> None:
    """Rejeita revisão obsoleta sem destruir a versão salva."""
    editor = Editor(Config(tmp_path))
    note = editor.current()
    editor.save(note["id"], note["title"], "Versão nova", note["revision"])
    with pytest.raises(RevisionConflict):
        editor.save(note["id"], note["title"], "Versão velha", note["revision"])
    assert editor.current()["markdown"] == "Versão nova"


def test_concurrent_saves_only_one_wins(tmp_path: Path) -> None:
    """Duas conexões não salvam sobre a mesma revisão."""
    config = Config(tmp_path)
    first, second = Editor(config), Editor(config)
    note = first.current()

    def attempt(editor: Editor, text: str) -> bool:
        try:
            editor.save(note["id"], note["title"], text, note["revision"])
            return True
        except RevisionConflict:
            return False

    with ThreadPoolExecutor() as pool:
        futures = [
            pool.submit(attempt, first, "A"),
            pool.submit(attempt, second, "B"),
        ]
        assert sorted(future.result() for future in futures) == [False, True]


def test_recovery_survives_restart(tmp_path: Path) -> None:
    """Recuperação persiste sem modificar a nota original."""
    config = Config(tmp_path)
    editor = Editor(config)
    note = editor.current()
    assert editor.stash_recovery(
        note["id"], "Rascunho", "Ainda não salvo", note["revision"]
    )["stored"]
    restored = Editor(config).current()
    assert restored["markdown"] == ""
    assert restored["recovery"]["markdown"] == "Ainda não salvo"
    editor.save(note["id"], "Rascunho", "Ainda não salvo", note["revision"])
    assert editor.current()["recovery"] is None
    assert not editor.stash_recovery(
        note["id"], "Velho", "Snapshot atrasado", note["revision"]
    )["stored"]


@pytest.mark.parametrize("identifier", ["../data", "invalid", "", None])
def test_invalid_identifier(tmp_path: Path, identifier: str) -> None:
    """Identificadores nunca são interpretados como caminhos."""
    with pytest.raises(ValueError):
        Editor(Config(tmp_path)).get(identifier)


def test_exports_preserve_text_and_structure(tmp_path: Path) -> None:
    """Lê arquivos exportados para verificar Unicode e estrutura."""
    pdf_path = tmp_path / "document.pdf"
    docx_path = tmp_path / "document.docx"
    export_pdf(MARKDOWN, pdf_path, "Ciência")
    export_docx(MARKDOWN, docx_path, "Ciência")
    pdf = PdfReader(pdf_path)
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "Ciência e ação" in text and "citação" in text
    document = Document(docx_path)
    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.tables[0].cell(1, 1).text == "42"
    runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
    assert any(run.bold and "forte" in run.text for run in runs)
    assert any(run.italic and "itálico" in run.text for run in runs)


def test_pdf_multipage_and_no_remote_images(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Pagina texto longo e impede busca de imagens na exportação."""
    import urllib.request

    def forbidden(*args: object, **kwargs: object) -> None:
        raise AssertionError("A exportação não deve acessar a rede.")

    monkeypatch.setattr(urllib.request, "urlopen", forbidden)
    target = tmp_path / "long.pdf"
    text = "Texto com acentuação e conteúdo de estudo.\n\n" * 180
    text += "![Exemplo](https://invalid.example/image.png)"
    export_pdf(text, target, "Longo")
    pdf = PdfReader(target)
    assert len(pdf.pages) > 1
    assert "[Imagem: Exemplo]" in pdf.pages[-1].extract_text()
