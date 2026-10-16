"""Original preservation, rotation geometry, native text and revision boundaries."""

import asyncio
import hashlib
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import pytest
from fpdf import FPDF
from pypdf import PdfReader, PdfWriter
from pptx import Presentation
from lxml import etree
from apps.open_doc.services.container import configure
from apps.open_doc.services.preferences_service import Config
from apps.open_doc.services.document_service import RevisionConflict
from apps.open_doc.ink.validation import validate_layers
from apps.open_doc.ink.handwriting import HandwritingService, Recognition
from apps.open_doc.document_model.schema import RichDocument, Paragraph, TextRun
from apps.open_doc.exporters.docx.package_writer import DocxPackageWriter


def stroke():
    return {
        "id": "s1",
        "tool": "pen",
        "color": "#152637",
        "size": 3,
        "opacity": 1,
        "points": [
            {
                "x": 0.1,
                "y": 0.1,
                "pressure": 0.2,
                "tilt_x": 10,
                "tilt_y": -10,
                "timestamp_ms": 1,
            },
            {"x": 0.3, "y": 0.2, "pressure": 0.9, "timestamp_ms": 2},
        ],
        "outline": [[0.1, 0.1], [0.3, 0.2], [0.3, 0.21], [0.1, 0.11]],
    }


def text():
    return {
        "id": "text1",
        "text": "Escrita convertida",
        "font_family": "Arial",
        "font_size_pt": 16,
        "color": "#152637",
        "x": 0.2,
        "y": 0.4,
        "width": 0.6,
        "height": 0.1,
        "confidence": 0.9,
    }


@pytest.fixture
def env(tmp_path):
    return configure(Config(tmp_path))


def pdf_source(root):
    source = root / "original.pdf"
    pdf = FPDF(unit="pt", format=(300, 400))
    for index in range(4):
        pdf.add_page()
        pdf.set_font("Helvetica", size=16)
        pdf.text(30, 40, "Original " + str(index))
    reader = PdfReader(BytesIO(pdf.output()))
    writer = PdfWriter()
    for index, page in enumerate(reader.pages):
        page.rotate(index * 90)
        writer.add_page(page)
    with source.open("wb") as stream:
        writer.write(stream)
    return source


def test_pdf_annotations_rotation_recovery_and_source_preserved(env, tmp_path):
    source = pdf_source(tmp_path)
    original = source.read_bytes()
    service = env.fixed_documents
    note = service.import_pdf(source)
    assert [(s["width_pt"], s["height_pt"]) for s in note["content"]["surfaces"]] == [
        (300, 400),
        (400, 300),
        (300, 400),
        (400, 300),
    ]
    model = deepcopy(note["content"])
    for index, surface in enumerate(model["surfaces"]):
        surface["strokes"] = [{**stroke(), "id": "s" + str(index)}]
        surface["text_overlays"] = [{**text(), "id": "t" + str(index)}]
    service.recovery(note["id"], note["title"], model, 1)
    assert service.get(note["id"])["recovery"]
    saved = service.save(note["id"], note["title"], model, 1)
    assert saved["revision"] == 2 and not saved["recovery"]
    with pytest.raises(RevisionConflict):
        service.save(note["id"], "stale", model, 1)
    with pytest.raises(RevisionConflict):
        service.recovery(note["id"], "stale", model, 1)
    dest = tmp_path / "annotated.pdf"
    service.export(note["id"], 2, "pdf", dest)
    reader = PdfReader(dest)
    assert len(reader.pages) == 4
    for page in reader.pages:
        assert "Escrita convertida" in page.extract_text()
        assert page.rotation == 0
        assert len(page.images) >= 1
    assert source.read_bytes() == original
    assert (
        hashlib.sha256(service.source(note["id"]).read_bytes()).hexdigest()
        == model["source_sha256"]
    )
    with pytest.raises(ValueError):
        service.export(note["id"], 2, "pdf", source)
    assert "source_path" not in note and str(tmp_path) not in str(note)
    changed = deepcopy(model)
    changed["surfaces"][0]["width_pt"] = 900
    with pytest.raises(ValueError):
        service.save(note["id"], "bad", changed, 2)
    assert service.duplicate(note["id"])["content"] == model


def test_pptx_native_text_and_transparent_full_slide_ink(env, tmp_path, monkeypatch):
    monkeypatch.setattr(
        "apps.open_doc.services.fixed_document_service.libreoffice_executable",
        lambda: None,
    )
    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    for _ in range(3):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    original = source.read_bytes()
    s = env.fixed_documents
    note = s.import_pptx(source)
    assert not note["capabilities"]["preview"]
    model = note["content"]
    model["surfaces"][0].update(strokes=[stroke()], text_overlays=[text()])
    note = s.save(note["id"], note["title"], model, 1)
    dest = tmp_path / "annotated.pptx"
    s.export(note["id"], 2, "pptx", dest)
    result = Presentation(dest)
    assert len(result.slides) == 3 and len(result.slides[0].shapes) == 2
    assert result.slides[0].shapes[1].text == "Escrita convertida"
    assert (
        result.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.name == "Arial"
    )
    assert source.read_bytes() == original


def test_docx_ink_is_anchored_and_source_opaque_parts_survive(tmp_path):
    model = RichDocument(
        blocks=[Paragraph(children=[TextRun(text="Texto nativo")])]
    ).to_dict()
    model["ink_layers"] = [
        {
            "surface_index": 0,
            "anchor_block_id": model["blocks"][0]["id"],
            "width_pt": 595,
            "height_pt": 842,
            "strokes": [stroke()],
        }
    ]
    source = tmp_path / "source.docx"
    DocxPackageWriter().write(model, source)
    with ZipFile(source, "a") as z:
        z.writestr("customXml/item99.xml", "<private>preservado</private>")
    target = tmp_path / "target.docx"
    DocxPackageWriter(source).write(model, target)
    with ZipFile(target) as z:
        assert z.read("customXml/item99.xml") == b"<private>preservado</private>"
        xml = etree.fromstring(z.read("word/document.xml"))
        anchors = xml.findall(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor"
        )
        assert anchors
        assert (
            anchors[0]
            .find(
                "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}positionH"
            )
            .get("relativeFrom")
            == "page"
        )
        assert any(name.startswith("word/media/opendoc-ink-") for name in z.namelist())


def test_mock_handwriting_is_async_and_missing_backend_keeps_strokes():
    class Fake:
        available = True
        name = "fake-local"

        async def recognize(self, strokes, *, language):
            return Recognition("Olá mundo", 0.92, self.name, language)

    strokes = [stroke()]
    original = deepcopy(strokes)
    assert (
        asyncio.run(HandwritingService([Fake()]).recognize(strokes)).text == "Olá mundo"
    )
    with pytest.raises(RuntimeError, match="Sua tinta foi mantida"):
        asyncio.run(HandwritingService([]).recognize(strokes))
    assert strokes == original


@pytest.mark.parametrize("bad", [float("nan"), -1, 1.1, "0.2", True])
def test_ink_rejects_invalid_geometry(bad):
    s = stroke()
    s["points"][0]["x"] = bad
    with pytest.raises(ValueError):
        validate_layers([{"surface_index": 0, "strokes": [s]}])


def test_artifact_http_requires_token_and_supports_range(tmp_path, monkeypatch):
    from webtest import TestApp
    from config.wsgi import create_app
    import vela.core.app as runtime
    from apps.open_doc.services.container import services

    class Window:
        def __init__(self, **kwargs):
            pass

        def run(self):
            pass

    monkeypatch.setattr(runtime, "DesktopWindow", Window)
    app = create_app(Config(tmp_path))
    note = services().fixed_documents.import_pdf(pdf_source(tmp_path))
    app.run()
    client = TestApp(app.api_server.app)
    url = "/api/open_doc/artifact/content?id=" + note["id"]
    client.get(url, status=403)
    headers = {"X-Open-Doc-Token": services().token}
    assert client.get(url, headers=headers).body.startswith(b"%PDF")
    response = client.get(url, headers={**headers, "Range": "bytes=0-15"}, status=206)
    assert len(response.body) == 16 and response.headers["Cache-Control"] == "no-store"
    client.get(
        "/api/open_doc/artifact/content?id=../../etc/passwd",
        headers=headers,
        status=404,
    )
    client.get(
        "/api/open_doc/artifact/content?path=/etc/passwd", headers=headers, status=404
    )


@pytest.mark.skipif(
    __import__("sys").platform != "win32"
    or __import__("os").environ.get("OPEN_DOC_TEST_WINDOWS_INK") != "1",
    reason="Windows Ink requires Windows and installed handwriting language",
)
def test_real_windows_ink_optional():
    from apps.open_doc.ink.windows_ink import WindowsInkBackend

    backend = WindowsInkBackend()
    if not backend.available:
        pytest.skip("Install requirements-windows-ink.txt and a handwriting language")
    result = asyncio.run(backend.recognize([stroke()], language="pt-BR"))
    assert result.backend == "windows-ink"


def test_rich_pdf_ink_reaches_approximate_page_export(env, tmp_path):
    model = RichDocument(
        blocks=[Paragraph(children=[TextRun(text="Rich PDF")])]
    ).to_dict()
    model["ink_layers"] = [
        {"surface_index": 0, "strokes": [stroke()], "width_pt": 595, "height_pt": 842}
    ]
    note = env.editor.current()
    note = env.editor.save(
        note["id"],
        "Ink PDF",
        "Rich PDF",
        note["revision"],
        content=model,
        content_format="opendoc-rich-v1",
    )
    target = tmp_path / "rich-ink.pdf"
    result = env.exports.export(note["id"], note["revision"], "pdf", target)
    assert len(PdfReader(target).pages[0].images) > 0
    assert any("Tinta" in warning for warning in result["warnings"])


def test_recovery_checks_revision_atomically_after_validation(
    env, tmp_path, monkeypatch
):
    service = env.fixed_documents
    note = service.import_pdf(pdf_source(tmp_path))
    original = service._checked

    def other_window_saves(identifier, content):
        old = original(identifier, content)
        monkeypatch.setattr(service, "_checked", original)
        service.save(identifier, "Outra janela", content, old["revision"])
        return old

    monkeypatch.setattr(service, "_checked", other_window_saves)
    with pytest.raises(RevisionConflict):
        service.recovery(note["id"], note["title"], note["content"], note["revision"])
    current = service.get(note["id"])
    assert current["title"] == "Outra janela" and current["recovery"] is None
