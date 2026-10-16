"""Real Qt input retention checks; no physical tablet or fake recognition claim."""

import os
import sys
import time
import tempfile
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
os.environ.setdefault("PYWEBVIEW_GUI", "qt")
os.environ.setdefault("QT_API", "pyqt5")
import webview
from config.wsgi import create_app
from apps.open_doc.services.container import services
from apps.open_doc.services.preferences_service import Config
from qa_runtime import Keyboard, report
from docx import Document
import subprocess


def wait(window, code):
    for _ in range(200):
        if window.evaluate_js(code):
            return
        time.sleep(0.1)
    raise AssertionError(code)


def main():
    errors = []
    original = webview.start

    def qa():
        window = webview.windows[0]
        try:
            report(window)
            wait(window, "document.querySelector('[data-app]')?.dataset.ready==='true'")
            if not os.environ.get("OPEN_DOC_INPUT_REOPEN"):
                window.evaluate_js("OpenDocApp.action('new')")
            if os.environ.get("OPEN_DOC_INPUT_NOTE_ID"):
                window.evaluate_js(
                    "OpenDocApp.editor().load("
                    + json.dumps(
                        services().editor.get(os.environ["OPEN_DOC_INPUT_NOTE_ID"])
                    )
                    + ")"
                )
            else:
                wait(
                    window,
                    "document.querySelector('[data-title]').value==='Documento sem título'",
                )
            time.sleep(0.3)
            window.evaluate_js(
                "(() => {const box=document.querySelector('[data-page-content]'),range=document.createRange();range.selectNodeContents(box);range.collapse(false);document.querySelector('[data-rich]').focus();getSelection().removeAllRanges();getSelection().addRange(range);})()"
            )
            keyboard = Keyboard(window)
            window.evaluate_js(
                "window.modelContainsText=(model,expected)=>{const text=v=>!v||typeof v!=='object'?'':typeof v.text==='string'?v.text:Object.values(v).map(text).join('');return text(model).includes(expected);}"
            )

            def assert_text(expected):
                if not window.evaluate_js(
                    "document.querySelector('[data-rich]').textContent.includes("
                    + json.dumps(expected)
                    + ")"
                ):
                    print(
                        "QA_FIXTURE_MISMATCH",
                        window.evaluate_js(
                            "({expected:"
                            + json.dumps(expected)
                            + ",matches:document.querySelector('[data-rich]').textContent.match(/Teste.{0,140}/g),model:OpenDocModel.markdown(OpenDocApp.editor().snapshot().content).slice(-180)})"
                        ),
                        flush=True,
                    )
                assert window.evaluate_js(
                    "document.querySelector('[data-rich]').textContent.includes("
                    + json.dumps(expected)
                    + ")"
                ), "DOM lost expected fixture text"
                assert window.evaluate_js(
                    "modelContainsText(OpenDocApp.editor().snapshot().content,"
                    + json.dumps(expected)
                    + ")"
                ), "Model lost expected fixture text"

            def retained(expected):
                for delay in [0.1, 0.3, 0.6, 1.2, 2.5]:
                    time.sleep(delay)
                    assert_text(expected)

            if os.environ.get("OPEN_DOC_INPUT_REOPEN"):
                for expected in json.loads(
                    os.environ.get(
                        "OPEN_DOC_INPUT_EXPECTED", '["Open Doc teclado 123"]'
                    )
                ):
                    assert_text(expected)
                print("QT_INPUT_REOPEN_SEPARATE_PROCESS_OK", flush=True)
                return
            keyboard.type("Open Doc teclado 123")
            retained("Open Doc teclado 123")
            print("QT_NEW_DOCUMENT_RETAINS_TYPED_TEXT_OK", flush=True)
            note_id = window.evaluate_js("OpenDocApp.editor().snapshot().id")
            deadline = time.monotonic() + 35
            while time.monotonic() < deadline:
                assert_text("Open Doc teclado 123")
                if "Open Doc teclado 123" in services().editor.get(note_id)["markdown"]:
                    break
                time.sleep(0.2)
            else:
                raise AssertionError("Autosave did not persist input")
            window.evaluate_js("OpenDocApp.editor().flush()")
            wait(
                window,
                "!document.querySelector('[data-app]').dataset.saving || document.querySelector('[data-app]').dataset.saving==='false'",
            )
            print("QT_INPUT_30_SECOND_AUTOSAVE_RETENTION_OK", flush=True)
            env = {**os.environ, "OPEN_DOC_INPUT_REOPEN": str(services().config.root)}
            reopened = subprocess.run(
                [sys.executable, __file__],
                env=env,
                capture_output=True,
                text=True,
                timeout=30,
            )
            assert reopened.returncode == 0, reopened.stdout + reopened.stderr
            print("QT_INPUT_REOPEN_SEPARATE_PROCESS_OK", flush=True)
            keyboard.key("Return")
            keyboard.type("Teste de digitação persistente 123 áéíóú ç.", delay=30)
            retained("Teste de digitação persistente 123 áéíóú ç.")
            keyboard.key("Space")
            keyboard.key("A", shift=True)
            keyboard.key("Backspace")
            keyboard.key("Delete")
            keyboard.press("B")
            keyboard.type("B")
            keyboard.press("B")
            assert window.evaluate_js(
                "OpenDocApp.editor().snapshot().content.blocks.length>=2"
            ), "Enter did not create a model paragraph"
            assert_text("Teste de digitação persistente 123 áéíóú ç.")
            print(
                "QT_INPUT_ENTER_BACKSPACE_DELETE_SPACE_SHIFT_ACCENTS_BOLD_OK",
                flush=True,
            )
            source = services().config.root / "large-input.docx"
            doc = Document()
            for index in range(20):
                doc.add_paragraph(
                    "Página "
                    + str(index + 1)
                    + " Texto de teste para entrada e histórico. " * 12
                )
                if index == 13:
                    doc.add_paragraph("Lista numerada QA", style="List Number")
                    doc.add_paragraph("Lista com marcador QA", style="List Bullet")
                if index == 14:
                    table = doc.add_table(rows=2, cols=2)
                    table.cell(0, 0).text = "Célula original"
                    from PIL import Image

                    image = services().config.root / "qa-image.png"
                    Image.new("RGB", (30, 30), "blue").save(image)
                    doc.add_picture(str(image))
                if index < 19:
                    doc.add_page_break()
            doc.add_page_break()
            doc.add_paragraph("fragmento " * 1800)
            doc.add_page_break()
            from docx.shared import Pt

            doc.add_paragraph().add_run("Overflow start ").font.size = Pt(48)
            doc.save(source)
            note = services().editor.import_docx(source)
            window.evaluate_js("OpenDocApp.editor().load(" + json.dumps(note) + ")")
            wait(window, "document.querySelectorAll('.document-page').length>=20")
            time.sleep(1)
            window.evaluate_js(
                "window.__fullRenders=0;const r=OpenDocRichRenderer.render;OpenDocRichRenderer.render=(...a)=>{if(a[1].matches('[data-rich]'))__fullRenders++;return r(...a);};window.__images=Array.from(document.querySelectorAll('[data-rich] img'));"
            )

            def caret(selector):
                window.evaluate_js(
                    """(() => {const p=document.querySelector(%s);if(!p)throw Error('QA target absent');p.scrollIntoView();const r=document.createRange();r.selectNodeContents(p);r.collapse(false);document.querySelector('[data-rich]').focus();getSelection().removeAllRanges();getSelection().addRange(r);OpenDocApp.editor().capture();})()"""
                    % json.dumps(selector)
                )

            count = window.evaluate_js(
                "document.querySelectorAll('.document-page').length"
            )
            for index in [0, 14, count - 1]:
                caret('[data-page-index="' + str(index) + '"] p')
                phrase = "Teste de digitação persistente 123 áéíóú ç. " + str(index)
                keyboard.type(phrase, delay=30)
                retained(phrase)
                top = window.evaluate_js(
                    "document.querySelector('[data-canvas]').scrollTop"
                )
                keyboard.press("Z")
                time.sleep(0.6)
                assert not window.evaluate_js(
                    "modelContainsText(OpenDocApp.editor().snapshot().content,"
                    + json.dumps(phrase)
                    + ")"
                )
                keyboard.press("Y")
                retained(phrase)
                assert (
                    abs(
                        window.evaluate_js(
                            "document.querySelector('[data-canvas]').scrollTop"
                        )
                        - top
                    )
                    < 4
                )
                assert window.evaluate_js(
                    "__fullRenders===0 && __images.every(i=>i.isConnected)"
                )
            print(
                "QT_INPUT_20_PLUS_PAGES_FIRST_15_LAST_RETAIN_UNDO_REDO_NO_FULL_RENDER_OK",
                flush=True,
            )
            for style in ["number", "bullet"]:
                selector = "[data-rich] [data-list-info]"
                candidates = window.evaluate_js(
                    "Array.from(document.querySelectorAll('[data-rich] [data-node-id]')).filter(n=>n.textContent.includes('Lista ')).map(n=>n.dataset.nodeId)"
                )
                assert candidates, "Imported list absent"
                caret(
                    '[data-node-id="'
                    + candidates[0 if style == "number" else -1]
                    + '"]'
                )
                phrase = " LISTA_QA_" + style
                keyboard.type(phrase, delay=30)
                retained(phrase)
                keyboard.press("Z")
                time.sleep(0.6)
                keyboard.press("Y")
                retained(phrase)
                assert window.evaluate_js(
                    "__fullRenders===0 && __images.every(i=>i.isConnected)"
                )
            print("QT_INPUT_NUMBERED_BULLET_LISTS_LOCAL_UNDO_OK", flush=True)
            caret("[data-rich] td p")
            keyboard.type(" Célula persistente áéíóú", delay=30)
            retained("Célula persistente áéíóú")
            keyboard.press("Z")
            time.sleep(0.6)
            keyboard.press("Y")
            retained("Célula persistente áéíóú")
            assert window.evaluate_js(
                "__fullRenders===0 && __images.every(i=>i.isConnected)"
            )
            caret("[data-rich] p[data-fragment-of]")
            keyboard.type(" FRAGMENTO_QA ", delay=450)
            retained("FRAGMENTO_QA")
            for _ in range(10):
                keyboard.type("x", delay=550)
            for _ in range(10):
                keyboard.press("Z")
            assert window.evaluate_js(
                "__fullRenders===0 && __images.every(i=>i.isConnected)"
            )
            print(
                "QT_INPUT_TABLE_FRAGMENT_SLOW_TYPING_10_UNDOS_IMAGE_STABILITY_OK",
                flush=True,
            )
            caret(
                '[data-page-index="'
                + str(
                    window.evaluate_js(
                        "document.querySelectorAll('.document-page').length-1"
                    )
                )
                + '"] p'
            )
            pages_before = window.evaluate_js(
                "document.querySelectorAll('.document-page').length"
            )
            keyboard.type("OVERFLOW_QA " * 40 + "FINAL_OVERFLOW_QA", delay=30)
            retained("FINAL_OVERFLOW_QA")
            assert (
                window.evaluate_js("document.querySelectorAll('.document-page').length")
                > pages_before
            )
            assert window.evaluate_js(
                "__fullRenders===0 && __images.every(i=>i.isConnected)"
            )
            print("QT_INPUT_OVERFLOW_NEW_PAGE_RETENTION_OK", flush=True)
            window.evaluate_js("OpenDocApp.editor().flush()")
            saved_id = window.evaluate_js("OpenDocApp.editor().snapshot().id")
            expected = [
                "Teste de digitação persistente 123 áéíóú ç. " + str(i)
                for i in [0, 14, count - 1]
            ] + [
                "Célula persistente áéíóú",
                "LISTA_QA_number",
                "LISTA_QA_bullet",
                "FINAL_OVERFLOW_QA",
            ]
            deadline = time.monotonic() + 15
            while (
                "FINAL_OVERFLOW_QA" not in services().editor.get(saved_id)["markdown"]
            ):
                assert time.monotonic() < deadline, "Large document save timeout"
                time.sleep(0.1)
            env = {
                **os.environ,
                "OPEN_DOC_INPUT_REOPEN": str(services().config.root),
                "OPEN_DOC_INPUT_NOTE_ID": saved_id,
                "OPEN_DOC_INPUT_EXPECTED": json.dumps(expected),
            }
            reopened = subprocess.run(
                [sys.executable, __file__],
                env=env,
                capture_output=True,
                text=True,
                timeout=30,
            )
            assert reopened.returncode == 0, reopened.stdout + reopened.stderr
            print("QT_INPUT_LARGE_DOCX_SAVE_REOPEN_ALL_EDITED_BLOCKS_OK", flush=True)
        except Exception as error:
            import traceback

            traceback.print_exc()
            errors.append(str(error))
        finally:
            window.confirm_close = False
            window.destroy()

    def start(*args, **kwargs):
        original(qa, gui="qt", debug=False)

    webview.start = start
    try:
        if os.environ.get("OPEN_DOC_INPUT_REOPEN"):
            app = create_app(Config(Path(os.environ["OPEN_DOC_INPUT_REOPEN"])))
            app.run()
        else:
            with tempfile.TemporaryDirectory(prefix="open-doc-input-") as folder:
                app = create_app(Config(Path(folder)))
                app.run()
    finally:
        webview.start = original
    if errors:
        raise RuntimeError("; ".join(errors))


if __name__ == "__main__":
    main()
