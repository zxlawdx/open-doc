"""Smoke test no Vela/Qt real e capturas via protocolo do QtWebEngine.

Executar em sessão gráfica; para CI configure QT_QPA_PLATFORM=offscreen.
Requer Node 22+ apenas para CDP/capturas. Não altera o workspace do usuário.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
os.environ.setdefault("PYWEBVIEW_GUI", "qt")
os.environ.setdefault("QT_API", "pyqt5")
os.environ.setdefault("QTWEBENGINE_REMOTE_DEBUGGING", "9234")

import webview
from apps.open_doc.services.preferences_service import Config
from apps.open_doc.services.container import services
from config.wsgi import create_app

FIXTURE = (
    "# Escrever com clareza\n\n"
    "Um documento começa com uma ideia. Este texto de validação verifica "
    "acentos, seleção e formatação no Open Doc.\n\n"
    "## Organização do texto\n\n"
    "A página ocupa o centro da mesa de trabalho. As ferramentas ficam "
    "por perto, sem disputar atenção com o conteúdo.\n\n"
    "- Registrar a primeira ideia\n- Desenvolver o argumento\n"
    "- Revisar com cuidado\n\n"
    "## Próximos passos\n\n"
    "O conteúdo permanece no dispositivo e pode ser exportado para PDF, "
    "DOCX ou Markdown.\n\n"
    "> Clareza também é escolher o que deixar de fora.\n"
)


def wait_for(window: Any, expression: str) -> Any:
    """Aguarda uma condição no renderer sem bloquear a thread Qt."""
    for _ in range(150):
        value = window.evaluate_js(expression)
        if value:
            return value
        time.sleep(0.1)
    raise AssertionError("Timeout: " + expression)


def native_keyboard(window: Any) -> None:
    """Envia atalhos Qt reais ao focusProxy do QWebEngineView."""
    from threading import Event
    from PyQt5.QtCore import QObject, Qt, pyqtSignal, pyqtSlot
    from PyQt5.QtTest import QTest
    from PyQt5.QtWidgets import QApplication
    from webview.platforms.qt import BrowserView

    class Driver(QObject):
        """Entrega eventos na thread gráfica do Qt sem acessar widgets em paralelo."""

        requested = pyqtSignal(str, bool)

        def __init__(self, view: Any) -> None:
            super().__init__()
            self.view = view
            self.done = Event()
            self.requested.connect(self.perform)

        @pyqtSlot(str, bool)
        def perform(self, key: str, shift: bool) -> None:
            target = self.view.focusProxy() or self.view
            target.setFocus()
            modifiers = Qt.ControlModifier | (
                Qt.ShiftModifier if shift else Qt.NoModifier
            )
            QTest.keyClick(target, getattr(Qt, "Key_" + key), modifiers)
            self.done.set()

        def send(self, key: str, shift: bool = False) -> None:
            self.done.clear()
            self.requested.emit(key, shift)
            assert self.done.wait(3)
            time.sleep(0.2)

    driver = Driver(BrowserView.instances[window.uid].webview)
    driver.moveToThread(QApplication.instance().thread())
    window.evaluate_js("document.querySelector('[data-rich]').focus()")
    time.sleep(0.2)
    driver.send("A")
    assert len(window.evaluate_js("window.getSelection().toString()")) > 50
    driver.send("B")
    assert "**" in window.evaluate_js("document.querySelector('[data-source]').value")
    driver.send("Z")
    driver.send("A")
    driver.send("I")
    assert window.evaluate_js(
        "Boolean(document.querySelector('[data-rich] i, [data-rich] em'))"
    )
    driver.send("Z")
    driver.send("A")
    driver.send("C")
    driver.send("X")
    assert not window.evaluate_js(
        "document.querySelector('[data-rich]').textContent"
    ).strip()
    driver.send("V")
    assert "Um documento começa" in window.evaluate_js(
        "document.querySelector('[data-rich]').textContent"
    )
    driver.send("Z")
    assert window.evaluate_js("document.querySelector('[data-source]').value").strip()
    driver.send("Y")
    driver.send("Z")
    driver.send("Z", shift=True)
    driver.send("S")
    wait_for(window, "window.OpenDocEditor.isClean()")
    print("NATIVE_CTRL_A_C_X_V_Z_Y_SHIFT_Z_S_OK", flush=True)


def verify_controls(window: Any) -> None:
    """Exercise UI handlers, selection and canonical output, never editor.format."""

    def js(expression: str) -> Any:
        return window.evaluate_js(expression)

    def click(selector: str) -> None:
        js("""(() => { const b=document.querySelector(%s);
          if(!b || b.disabled) throw Error('Missing/disabled control: '+%s);
          b.dispatchEvent(new MouseEvent('mousedown',{bubbles:true,cancelable:true}));
          b.click(); })()""" % (json.dumps(selector), json.dumps(selector)))
        time.sleep(0.15)

    def select() -> None:
        # Reset is test setup; actions below always go through visible controls.
        note = services().editor.current()
        note = {**note, "markdown": "# Título\n\nTexto selecionado"}
        js("window.OpenDocApp.editor().load(" + json.dumps(note) + ")")
        js("""(() => {const r=document.querySelector('[data-rich]');r.focus();
          const range=document.createRange();range.selectNodeContents(r.querySelector('p'));
          const s=getSelection();s.removeAllRanges();s.addRange(range);})()""")

    for command, expected in [
        ("bold", "**Texto selecionado**"),
        ("italic", "_Texto selecionado_"),
        ("list", "- Texto"),
        ("ordered", "1. Texto"),
        ("quote", "> Texto"),
        ("code", "```\nTexto selecionado\n```"),
    ]:
        select()
        click(".toolbar [data-command=" + command + "]")
        value = js("document.querySelector('[data-source]').value")
        assert expected in value, (command, value)
        if command == "code":
            assert js("Boolean(document.querySelector('[data-rich] pre'))")
    for style, tag in [
        ("h1", "h1"),
        ("heading", "h2"),
        ("h3", "h3"),
        ("paragraph", "p"),
    ]:
        select()
        js(
            """(() => {const s=document.querySelector('[data-paragraph]');
          s.dispatchEvent(new MouseEvent('mousedown',{bubbles:true}));s.focus();
          s.value=%s;s.dispatchEvent(new Event('change',{bubbles:true}));})()"""
            % json.dumps(style)
        )
        assert (
            js(
                "document.querySelector('[data-rich]').lastElementChild.tagName.toLowerCase()"
            )
            == tag
        )
    select()
    click(".toolbar [data-action=link]")
    js("document.querySelector('[data-dialog-input]').value='https://example.com'")
    click("[data-dialog-confirm]")
    assert "[Texto selecionado](https://example.com)" in js(
        "document.querySelector('[data-source]').value"
    ), js(
        "JSON.stringify({source:document.querySelector('[data-source]').value,html:document.querySelector('[data-rich]').innerHTML})"
    )
    select()
    click(".toolbar [data-action=link]")
    click("[data-dialog-cancel]")
    assert "](" not in js("document.querySelector('[data-source]').value")
    select()
    click(".toolbar [data-action=link]")
    js("document.querySelector('[data-dialog-input]').value='javascript:alert(1)'")
    click("[data-dialog-confirm]")
    assert "](" not in js("document.querySelector('[data-source]').value")
    for menu in ["file", "edit", "view", "insert", "format", "tools", "help"]:
        click("[data-menu=" + menu + "]")
        assert (
            js("document.querySelector('[data-menu-panel=" + menu + "]').hidden")
            is False
        )
        click("[data-menu=" + menu + "]")
    for mode in ["markdown", "preview", "rich"]:
        click("[data-menu=view]")
        click("[data-mode=" + mode + "]")
        assert js("document.querySelector('[data-surface]').dataset.editorMode") == mode
    for command in [
        "bold",
        "italic",
        "paragraph",
        "h1",
        "heading",
        "h3",
        "list",
        "ordered",
        "quote",
        "code",
        "selectAll",
    ]:
        select()
        menu = (
            "edit"
            if command == "selectAll"
            else (
                "insert"
                if command in ["list", "ordered", "quote", "code"]
                else "format"
            )
        )
        click("[data-menu=" + menu + "]")
        click("[data-menu-panel=" + menu + "] [data-command=" + command + "]")
        if command == "selectAll":
            assert len(js("getSelection().toString()")) > 10
        else:
            assert (
                js("document.querySelector('[data-source]').value")
                != "# Título\n\nTexto selecionado"
                or command == "paragraph"
            ), (command, js("document.querySelector('[data-rich]').innerHTML"))
    select()
    click("[data-menu=tools]")
    click("[data-menu-panel=tools] [data-action=find]")
    js(
        "document.querySelector('[data-find]').value='Texto';document.querySelector('[data-find]').dispatchEvent(new Event('input',{bubbles:true}))"
    )
    wait_for(window, "document.querySelectorAll('[data-find-results] button').length")
    click("[data-find-results] button")
    assert js("getSelection().toString()") == "Texto", js(
        "JSON.stringify({selected:getSelection().toString(),html:document.querySelector('[data-rich]').innerHTML,mark:window.OpenDocApp.editor().snapshot().selection})"
    )
    click("[data-outline] button")
    click("[data-menu=view]")
    click("[data-menu-panel=view] [data-action=navigator]")
    assert js("document.querySelector('[data-app]').dataset.collapsed") == "true"
    click("[data-menu=view]")
    click("[data-menu-panel=view] [data-action=navigator]")
    click('[data-zoom-step="10"]')
    assert js("document.querySelector('[data-zoom-value]').textContent") == "110%"
    click("[data-menu=view]")
    click("[data-action=zoom-reset]")
    assert js("document.querySelector('[data-zoom-value]').textContent") == "100%"
    click("[data-menu=tools]")
    click("[data-action=theme]")
    js("document.querySelector('[data-dialog-select]').value='dark'")
    click("[data-dialog-confirm]")
    assert js("document.querySelector('[data-app]').dataset.theme") == "dark"
    for action in ["shortcuts", "about"]:
        click("[data-menu=help]")
        click("[data-action=" + action + "]")
        assert js("document.querySelector('[data-dialog]').hidden") is False
        click("[data-dialog-confirm]")
    assert not js(
        "/is not a function|undefined|TypeError/.test(document.querySelector('[data-status]').textContent)"
    )
    assert js("window.OpenDocApp.failures.length") == 0
    print("NATIVE_TOOLBAR_PARAGRAPHS_LINK_MENUS_FIND_SIDEBAR_ZOOM_OK", flush=True)


def native_type(window: Any, text: str) -> None:
    """Type through actual Qt input events on its GUI thread."""
    from threading import Event
    from PyQt5.QtCore import QObject, pyqtSignal, pyqtSlot
    from PyQt5.QtTest import QTest
    from PyQt5.QtWidgets import QApplication
    from webview.platforms.qt import BrowserView

    class Typist(QObject):
        requested = pyqtSignal(str)

        def __init__(self):
            super().__init__()
            self.done = Event()
            self.requested.connect(self.perform)

        @pyqtSlot(str)
        def perform(self, value):
            view = BrowserView.instances[window.uid].webview
            target = view.focusProxy() or view
            target.setFocus()
            QTest.keyClicks(target, value)
            self.done.set()

    driver = Typist()
    driver.moveToThread(QApplication.instance().thread())
    driver.requested.emit(text)
    assert driver.done.wait(5)
    time.sleep(0.2)


def verify_rich_document(window: Any, source: Path) -> None:
    """Real app/bridge/API; only replace the file-picker choice in this QA."""
    from docx import Document
    from apps.open_doc.services.document_service import Editor

    def click(selector):
        window.evaluate_js(
            """(() => {const e=document.querySelector(%s);if(!e || e.disabled)throw Error('Control unavailable');e.dispatchEvent(new MouseEvent('mousedown',{bubbles:true,cancelable:true}));e.click();})()"""
            % json.dumps(selector)
        )
        time.sleep(0.15)

    def import_ui(path):
        original = window.create_file_dialog
        window.create_file_dialog = lambda *a, **kw: (str(path),)
        try:
            click("[data-menu=file]")
            click("[data-menu-panel=file] [data-action=import]")
            wait_for(
                window,
                "document.querySelector('[data-title]').value === "
                + json.dumps(path.stem),
            )
            wait_for(
                window,
                "window.OpenDocApp.editor().snapshot().content_format === 'opendoc-rich-v1'",
            )
        finally:
            window.create_file_dialog = original

    def export_ui():
        click("[data-menu=file]")
        click("[data-export=docx]")
        wait_for(
            window,
            "document.querySelector('[data-export-result]').textContent.startsWith('Salvo em: ')",
        )
        return Path(
            window.evaluate_js(
                "document.querySelector('[data-export-result]').textContent"
            ).removeprefix("Salvo em: ")
        )

    import_ui(source)
    assert (
        window.evaluate_js("document.querySelectorAll('[data-rich] table').length") > 0
    )
    window.evaluate_js(
        """(() => {const rich=document.querySelector('[data-rich]');rich.focus();const text=rich.querySelector('p [data-node-type=text]').firstChild;const range=document.createRange();range.selectNodeContents(text);getSelection().removeAllRanges();getSelection().addRange(range);})()"""
    )
    for selector, value in [
        ("[data-font-family]", "DejaVu Serif"),
        ("[data-font-size]", "18"),
    ]:
        window.evaluate_js(
            """(() => {const e=document.querySelector(%s);e.dispatchEvent(new MouseEvent('mousedown',{bubbles:true,cancelable:true}));e.focus();e.value=%s;e.dispatchEvent(new Event('change',{bubbles:true}));})()"""
            % (json.dumps(selector), json.dumps(value))
        )
    click(".toolbar [data-command=bold]")
    click(".toolbar [data-command=underline]")
    window.evaluate_js(
        """(() => {const e=document.querySelector('[data-rich] td [data-node-type=text]'),rich=e.closest('[data-rich]');rich.focus();const range=document.createRange();range.selectNodeContents(e);getSelection().removeAllRanges();getSelection().addRange(range);})()"""
    )
    native_type(window, "Celula editada no Qt")
    click("[data-menu=file]")
    click("[data-menu-panel=file] [data-action=save]")
    wait_for(window, "window.OpenDocEditor.isClean()")
    saved = services().editor.current()
    assert saved[
        "content_format"
    ] == "opendoc-rich-v1" and "Celula editada no Qt" in str(saved["content"])
    assert any(
        r.get("style", {}).get("font_family") == "DejaVu Serif"
        and r.get("style", {}).get("font_size_pt") == 18
        for b in saved["content"]["blocks"]
        for r in b.get("children", [])
    )
    exported = export_ui()
    doc = Document(exported)
    assert "Celula editada no Qt" in doc.tables[0].cell(0, 0).text
    assert any(
        r.font.name == "DejaVu Serif" and r.font.size.pt == 18
        for p in doc.paragraphs
        for r in p.runs
        if r.font.size
    )
    assert Editor(services().config).current()["content"] == saved["content"]
    print("NATIVE_RICH_DOCX_IMPORT_FORMAT_TABLE_SAVE_EXPORT_OK", flush=True)
    captured = subprocess.run(
        [
            "node",
            str(ROOT / "tools/qa_rich_capture.cjs"),
            os.environ["QTWEBENGINE_REMOTE_DEBUGGING"],
        ],
        text=True,
        capture_output=True,
        timeout=30,
    )
    print(captured.stdout, flush=True)
    assert captured.returncode == 0, captured.stderr
    user_source = os.environ.get("OPEN_DOC_USER_FIXTURE")
    if user_source:
        import_ui(Path(user_source))
        actual = services().editor.current()
        expected = Document(user_source)
        assert actual["content"]["footers"]
        assert (
            actual["content"]["sections"][0]["page_width_pt"]
            == expected.sections[0].page_width.pt
        )
        result = Document(export_ui())
        assert [p.text for p in expected.paragraphs] == [
            p.text for p in result.paragraphs
        ]
        print("NATIVE_USER_DOCX_IMPORT_FOOTER_FONTS_ROUNDTRIP_OK", flush=True)
    env = {
        **os.environ,
        "OPEN_DOC_QA_REOPEN": str(services().config.root),
        "OPEN_DOC_EXPECTED_ID": services().editor.current()["id"],
        "QTWEBENGINE_REMOTE_DEBUGGING": "9235",
    }
    reopened = subprocess.run(
        [sys.executable, str(Path(__file__))],
        env=env,
        text=True,
        capture_output=True,
        timeout=30,
    )
    print(reopened.stdout, flush=True)
    assert reopened.returncode == 0, reopened.stderr


def run_qa() -> None:
    """Executa VelaApp e verifica comandos no Chromium nativo do Qt."""
    original_start = webview.start
    failures: list[str] = []

    def qa() -> None:
        window = webview.windows[0]
        picker = window.create_file_dialog

        def chosen_file(dialog_type, *args, **kwargs):
            if dialog_type == webview.FileDialog.SAVE:
                return (
                    str(
                        services().config.root
                        / "exports"
                        / kwargs.get("save_filename", "qa.docx")
                    ),
                )
            return picker(dialog_type, *args, **kwargs)

        window.create_file_dialog = chosen_file
        try:
            wait_for(
                window,
                "Boolean(window.OpenDocApp && document.querySelector('[data-app]').dataset.ready === 'true')",
            )
            if os.environ.get("OPEN_DOC_QA_REOPEN"):
                snapshot = window.evaluate_js("window.OpenDocApp.editor().snapshot()")
                assert (
                    snapshot["id"] == os.environ["OPEN_DOC_EXPECTED_ID"]
                    and snapshot["content_format"] == "opendoc-rich-v1"
                )
                assert (
                    window.evaluate_js(
                        "document.querySelector('[data-rich]').textContent.length"
                    )
                    > 20
                )
                print("NATIVE_RICH_REOPEN_SEPARATE_PROCESS_OK", flush=True)
                return
            assert window.title == "Open Doc"
            assert (
                window.evaluate_js("document.querySelector('[data-rich]').hidden")
                is False
            )
            assert (
                window.evaluate_js("document.querySelector('[data-source]').hidden")
                is True
            )
            print("NATIVE_VELA_BOOT_OK", flush=True)
            assert window.evaluate_js(
                "(() => { const t = new EventTarget(); const c = new AbortController(); let n=0; t.addEventListener('x',()=>n++,{signal:c.signal}); c.abort(); t.dispatchEvent(new Event('x')); return n===0; })()"
            )
            window.evaluate_js(
                "window.__previousOpenDoc = window.OpenDocApp; Vela.navigate('/')"
            )
            wait_for(
                window,
                "Boolean(window.OpenDocApp !== window.__previousOpenDoc && document.querySelector('[data-app]').dataset.ready === 'true')",
            )
            print("NATIVE_VELA_REMOUNT_AND_LISTENER_CLEANUP_OK", flush=True)
            result = window.evaluate_js("""(() => {
              const rich = document.querySelector('[data-rich]');
              const text = rich.querySelector('p').firstChild;
              const range = document.createRange();range.setStart(text,0);range.setEnd(text,12);
              const selection=window.getSelection();rich.focus();selection.removeAllRanges();selection.addRange(range);
              const button=document.querySelector('.toolbar [data-command=bold]');
              button.dispatchEvent(new MouseEvent('mousedown',{bubbles:true,cancelable:true}));button.click();
              return {markdown:document.querySelector('[data-source]').value,selected:selection.toString()};
            })()""")
            assert "**" in result["markdown"] and result["selected"]
            window.evaluate_js(
                "document.querySelector('.toolbar [data-command=undo]').click()"
            )
            assert "**" not in window.evaluate_js(
                "document.querySelector('[data-source]').value"
            )
            window.evaluate_js(
                "document.querySelector('.toolbar [data-command=redo]').click()"
            )
            assert "**" in window.evaluate_js(
                "document.querySelector('[data-source]').value"
            )
            window.evaluate_js(
                "document.querySelector('.toolbar [data-command=italic]').click()"
            )
            assert "_" in window.evaluate_js(
                "document.querySelector('[data-source]').value"
            )
            print("NATIVE_SELECTION_BOLD_ITALIC_UNDO_REDO_OK", flush=True)
            verify_controls(window)
            sys.path.insert(0, str(ROOT / "tests"))
            from docx_fixtures import fixture_docx

            verify_rich_document(
                window, fixture_docx(services().config.root / "qa-rich.docx")
            )
            from qa_pages import verify_pages

            verify_pages(window, services, wait_for)
            fresh = services().editor.create("Documento de validação", FIXTURE)
            window.evaluate_js(
                "window.OpenDocApp.editor().load(" + json.dumps(fresh) + ")"
            )
            window.evaluate_js("window.OpenDocApp.editor().flush()")
            wait_for(window, "window.OpenDocEditor.isClean()")
            assert not window.confirm_close
            for extension in ("pdf", "docx", "md"):
                window.evaluate_js(
                    "window.__qaExport = null; window.OpenDocApp.editor().export('"
                    + extension
                    + "').then(result => window.__qaExport = result).catch(error => window.__qaExport = {error: error.message})"
                )
                exported = wait_for(window, "window.__qaExport")
                assert "error" not in exported, exported
                target = Path(exported["path"])
                assert target.suffix == "." + extension and target.stat().st_size > 0
            print("NATIVE_SAVE_EXPORT_OK", flush=True)
            keyboard_document = services().editor.create(
                "Documento de validação", FIXTURE
            )
            window.evaluate_js(
                "window.OpenDocApp.editor().load(" + json.dumps(keyboard_document) + ")"
            )
            native_keyboard(window)
            # Return the fixture to its unformatted original content for screenshots.
            document = services().editor.create("Documento de validação", FIXTURE)
            window.evaluate_js(
                "window.OpenDocApp.editor().load(" + json.dumps(document) + ")"
            )
            window.evaluate_js(
                "document.querySelector('[data-export-result]').textContent=''"
            )
            screenshot_script = ROOT / "tools" / "qa_capture.cjs"
            output = subprocess.run(
                [
                    "node",
                    str(screenshot_script),
                    os.environ["QTWEBENGINE_REMOTE_DEBUGGING"],
                ],
                text=True,
                capture_output=True,
                timeout=50,
            )
            print(output.stdout, output.stderr, flush=True)
            assert output.returncode == 0
            window.evaluate_js("window.OpenDocApp.showHome()")
            wait_for(
                window, "document.querySelectorAll('[data-documents] li').length > 0"
            )
            window.evaluate_js("window.OpenDocApp.action('new')")
            wait_for(
                window,
                "document.querySelector('[data-title]').value === 'Documento sem título'",
            )
            window.evaluate_js("window.OpenDocApp.editor().flush()")
            print("NATIVE_RECENT_NEW_AND_CLOSE_OK", flush=True)
        except Exception as error:
            import traceback

            traceback.print_exc()
            try:
                print(
                    "QA_STATUS",
                    window.evaluate_js(
                        "document.querySelector('[data-status]').textContent"
                    ),
                    flush=True,
                )
                (ROOT.parent / "rich-qa-failure.json").write_text(
                    json.dumps(
                        window.evaluate_js("window.OpenDocApp.editor().snapshot()")
                    ),
                    encoding="utf-8",
                )
            except Exception:
                pass
            failures.append(repr(error))
            print("NATIVE_QA_ERROR", repr(error), flush=True)
        finally:
            # Test cleanup only; avoid a modal blocking CI if an assertion failed.
            window.confirm_close = False
            window.destroy()

    def start(*args: Any, **kwargs: Any) -> None:
        original_start(qa, gui="qt", debug=False)

    webview.start = start
    try:
        from contextlib import nullcontext

        reopen = os.environ.get("OPEN_DOC_QA_REOPEN")
        with (
            nullcontext(reopen)
            if reopen
            else tempfile.TemporaryDirectory(prefix="open-doc-qa-")
        ) as folder:
            application = create_app(Config(Path(folder)))
            if not reopen:
                note = services().editor.current()
                services().editor.save(
                    note["id"], "Documento de validação", FIXTURE, note["revision"]
                )
            application.run()
    finally:
        webview.start = original_start
    if failures:
        raise RuntimeError("; ".join(failures))
    print("NATIVE_QA_COMPLETE", flush=True)


if __name__ == "__main__":
    run_qa()
